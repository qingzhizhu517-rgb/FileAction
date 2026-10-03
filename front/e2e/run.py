"""运行浏览器全栈流程测试。模型为明确标注的HTTP替身，数据在临时目录。"""
import os
from pathlib import Path
import shutil
import socket
import subprocess
import sys
import tempfile
import time
import urllib.request

ROOT = Path(__file__).resolve().parents[2]


def create_synthetic_documents(directory):
    from docx import Document
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    directory.mkdir()
    document = Document()
    document.add_paragraph('合成课程通知：只用于浏览器测试')
    document.add_table(rows=1, cols=1).cell(0, 0).text = '合成表格：请核实申请要求'
    document.save(directory / 'synthetic.docx')
    writer = PdfWriter()
    page = writer.add_blank_page(width=300, height=300)
    font = DictionaryObject({NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type1'), NameObject('/BaseFont'): NameObject('/Helvetica')})
    page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): writer._add_object(font)})})
    stream = DecodedStreamObject()
    stream.set_data(b'BT /F1 12 Tf 10 100 Td (Synthetic notice: browser test only.) Tj ET')
    page[NameObject('/Contents')] = writer._add_object(stream)
    with (directory / 'synthetic.pdf').open('wb') as output:
        writer.write(output)


def wait_ready(url, process):
    for _ in range(100):
        if process.poll() is not None:
            raise RuntimeError('测试服务提前退出')
        try:
            with urllib.request.urlopen(url, timeout=1) as response:
                if response.status == 200:
                    return
        except OSError:
            time.sleep(0.1)
    raise RuntimeError('本机测试服务未就绪')


def main():
    # 不接管或杀死用户已有监听服务。
    for port in (8776, 8777, 8778):
        with socket.socket() as probe:
            if probe.connect_ex(('127.0.0.1', port)) == 0:
                raise RuntimeError(f'测试端口 {port} 已被占用，请先停止此前测试服务')
    processes = []
    with tempfile.TemporaryDirectory(prefix='fileaction-e2e-') as temporary:
        fixture_directory = Path(temporary) / 'synthetic-fixtures'
        create_synthetic_documents(fixture_directory)
        env = dict(os.environ)
        env.update({
            'PYTHONPATH': str(ROOT / 'backend'),
            'FILEACTION_MODEL_BASE_URL': 'http://127.0.0.1:8777/v1',
            'FILEACTION_MODEL_API_KEY': 'synthetic-e2e-placeholder',
            'FILEACTION_MODEL_NAME': 'explicit-http-test-double',
            'E2E_URL': 'http://127.0.0.1:8776',
            'E2E_UNCONFIGURED_URL': 'http://127.0.0.1:8778',
            'E2E_FIXTURES_DIR': str(fixture_directory),
        })
        try:
            stub = subprocess.Popen([sys.executable, str(ROOT/'front/e2e/model_stub.py')], cwd=temporary, env=env, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
            processes.append(stub)
            wait_ready('http://127.0.0.1:8777', stub)
            server = subprocess.Popen([sys.executable, '-m', 'fileaction.legacy', '--no-env-file', '--port', '8776'], cwd=temporary, env=env, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
            processes.append(server)
            wait_ready('http://127.0.0.1:8776/api/config', server)
            unconfigured_env = {key: value for key, value in env.items() if not key.startswith('FILEACTION_MODEL_')}
            unconfigured_directory = Path(temporary) / 'unconfigured'
            unconfigured_directory.mkdir()
            unconfigured = subprocess.Popen([sys.executable, '-m', 'fileaction.legacy', '--no-env-file', '--port', '8778'], cwd=unconfigured_directory, env=unconfigured_env, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
            processes.append(unconfigured)
            wait_ready('http://127.0.0.1:8778/api/config', unconfigured)
            executable = shutil.which('node')
            playwright_cli = ROOT / 'front/node_modules/@playwright/test/cli.js'
            if not executable or not playwright_cli.is_file():
                raise RuntimeError('缺少Node或Playwright，请先安装开发依赖')
            # 直接调用Node，避免Windows的.cmd把正则参数中的管道符解释成shell语法。
            return subprocess.run([executable, str(playwright_cli), 'test', *sys.argv[1:]], cwd=ROOT / 'front', env=env).returncode
        finally:
            for process in reversed(processes):
                process.terminate()
                try:
                    process.wait(timeout=10)
                except subprocess.TimeoutExpired:
                    process.kill()
                    process.wait()


if __name__ == '__main__':
    raise SystemExit(main())
