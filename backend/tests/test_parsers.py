import io
import zipfile

import pytest

from fileaction.parsers import ParseError, parse_document


def test_line_locator_and_utf8_text():
    document = parse_document("通知.txt", "合成通知\n截止：10月15日".encode("utf-8"))
    assert document.name == "通知.txt"
    assert document.segments[0].locator == "第1行"
    assert document.segments[1].locator == "第2行"
    assert document.segments[1].text == "截止：10月15日"
    assert document.hash


def test_empty_and_invalid_utf8_are_rejected():
    with pytest.raises(ParseError, match="空文件"):
        parse_document("empty.txt", b"")
    with pytest.raises(ParseError, match="UTF-8"):
        parse_document("bad.txt", b"\xff\xfe")


def test_limits_reject_oversized_text_without_truncation():
    with pytest.raises(ParseError, match="1,000,000"):
        parse_document("large.md", ("x" * 1000001).encode())
    with pytest.raises(ParseError, match="10 MiB"):
        parse_document("large.txt", b"x" * (10 * 1024 * 1024 + 1))


def test_docx_extracts_paragraph_and_table_locations():
    from docx import Document as DocxDocument

    source = DocxDocument()
    source.add_paragraph("第一段")
    table = source.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "表格内容"
    stream = io.BytesIO()
    source.save(stream)
    document = parse_document("材料.docx", stream.getvalue())
    assert any(segment.locator == "第1段" and segment.text == "第一段" for segment in document.segments)
    assert any("表格" in segment.locator and segment.text == "表格内容" for segment in document.segments)


def test_zip_bomb_docx_is_rejected():
    payload = io.BytesIO()
    with zipfile.ZipFile(payload, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", b"x" * (20 * 1024 * 1024))
    with pytest.raises(ParseError, match="解压"):
        parse_document("bomb.docx", payload.getvalue())


def test_pdf_scanned_limit_corrupt_and_encrypted_are_rejected():
    from pypdf import PdfWriter
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    output = io.BytesIO()
    writer.write(output)
    with pytest.raises(ParseError, match="文本"):
        parse_document("扫描.pdf", output.getvalue())
    for _ in range(100):
        writer.add_blank_page(width=100, height=100)
    output = io.BytesIO()
    writer.write(output)
    with pytest.raises(ParseError, match="100"):
        parse_document("long.pdf", output.getvalue())
    writer.encrypt("test")
    output = io.BytesIO()
    writer.write(output)
    with pytest.raises(ParseError, match="加密"):
        parse_document("encrypted.pdf", output.getvalue())
    with pytest.raises(ParseError):
        parse_document("fake.pdf", b"not a pdf")


def test_binary_content_disguised_as_text_is_rejected():
    with pytest.raises(ParseError, match="文本"):
        parse_document("binary.txt", b"PK\x00\x00binary")


def test_text_whitespace_lines_are_skipped_preserving_locator():
    document = parse_document("space.txt", b"First line\n  \nLast line")
    assert [segment.locator for segment in document.segments] == ["第1行", "第3行"]


def test_repeated_tiny_lines_rejected_before_creating_segments(monkeypatch):
    import fileaction.parsers as parsers
    calls = 0
    original = parsers.Segment
    def counted_segment(**kwargs):
        nonlocal calls
        calls += 1
        return original(**kwargs)
    monkeypatch.setattr(parsers, "Segment", counted_segment)
    with pytest.raises(ParseError):
        parse_document("many.txt", b"x\n" * 1000001)
    assert calls <= 1000000


def test_text_pdf_extracts_real_page_text():
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
    writer = PdfWriter()
    page = writer.add_blank_page(width=100, height=100)
    font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica")})
    page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 10 10 Td (Synthetic notice) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(stream)
    output = io.BytesIO()
    writer.write(output)
    document = parse_document("synthetic.pdf", output.getvalue())
    assert document.segments[0].text == "Synthetic notice"
    assert document.segments[0].locator.startswith("第1页")


